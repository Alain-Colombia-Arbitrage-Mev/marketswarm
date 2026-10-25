#!/usr/bin/env python3
"""Generate a Word report for the Mantralia UX audit, for later PDF export."""
import json
import os
from datetime import datetime, timezone

from docx import Document
from docx.shared import Pt
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AUDIT = os.path.join(BASE, "output", "audit")
report = json.load(open(os.path.join(AUDIT, "report.json"), encoding="utf-8"))
stats = report["stats"]
screens = report["screens"]

def sort_key(r):
    c = r.get("clarity_0_10")
    return (c is None, c if isinstance(c, (int, float)) else 0)
screens = sorted(screens, key=sort_key)

# Swarm synthesis: pick the report with the most findings (the completed run).
swarm = None
for name in os.listdir(AUDIT):
    p = os.path.join(AUDIT, name, "report.json")
    if os.path.exists(p):
        cand = json.load(open(p, encoding="utf-8"))
        out = cand.get("synthesis", {}).get("output") or {}
        if swarm is None or len(out.get("findings", [])) > len(
                (swarm.get("synthesis", {}).get("output") or {}).get("findings", [])):
            swarm = cand
syn = (swarm or {}).get("synthesis", {}).get("output") or {}
syn_summary = syn.get("summary", "")
findings = [f.get("statement", "") for f in syn.get("findings", [])]
experiments = syn.get("next_experiments", [])
gaps = syn.get("gaps", [])

doc = Document()

# Landscape for the wide table.
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width

t = doc.add_heading("Mantralia — Auditoría UX", level=0)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
now = datetime.now(timezone.utc).strftime("%Y-%m-%d")
p.add_run(f"{now} · {stats.get('total_screens', len(screens))} pantallas · "
          "pen.dev + market-swarm (swarm de 12 agentes) + visión/OCR (Nemotron VL)").italic = True

doc.add_heading("Resumen ejecutivo", level=1)
summ = [
    ("Pantallas analizadas", str(stats.get("total_screens", len(screens)))),
    ("Claridad promedio (0-10)", f"{stats.get('avg_clarity')} / 10"),
    ("Pantallas con CTA claro", str(stats.get("with_ctas"))),
    ("Pantallas con problemas detectados", str(stats.get("with_issues"))),
    ("Pantallas que mencionan precio/plan", str(stats.get("with_price"))),
]
table = doc.add_table(rows=len(summ), cols=2)
table.style = "Table Grid"
for i, (k, v) in enumerate(summ):
    table.rows[i].cells[0].text = k
    table.rows[i].cells[1].text = v
    for cell in table.rows[i].cells:
        for par in cell.paragraphs:
            for run in par.runs:
                run.font.size = Pt(10)

doc.add_heading("Veredicto del swarm", level=1)
doc.add_paragraph(syn_summary)

doc.add_heading("Hallazgos clave", level=1)
for f in findings:
    doc.add_paragraph(f, style="List Bullet")

doc.add_heading("Las 10 pantallas con menor claridad", level=1)
worst = [s for s in screens if isinstance(s.get("clarity_0_10"), (int, float))][:10]
for s in worst:
    doc.add_paragraph(f"{s['clarity_0_10']}/10 — {s['name']}", style="List Bullet")

doc.add_heading("Próximos experimentos recomendados", level=1)
for e in experiments:
    doc.add_paragraph(e, style="List Bullet")

doc.add_heading("Datos que faltan (gaps)", level=1)
for g in gaps:
    doc.add_paragraph(g, style="List Bullet")


# ---- Pricing / Gabor-Granger (sintético) ----
doc.add_heading("Pricing — Estudio Gabor-Granger (sintético)", level=1)
doc.add_paragraph(
    "Piloto sintético con paneles artificiales (no es validación de demanda real). "
    "El precio más alto pagable se define como el mayor precio con aceptación ≥ 30%.")

GABOR = os.path.join(os.path.dirname(AUDIT), "gabor")
cs = json.load(open(os.path.join(GABOR, "countries-summary.json"), encoding="utf-8"))
sym = {"USD": "$", "MXN": "MX$", "BRL": "R$"}
ptbl = doc.add_table(rows=1, cols=4)
ptbl.style = "Table Grid"
for i, h in enumerate(["País", "Moneda", "Precio más alto pagable (≥30%)",
                       "Precio óptimo por ingreso"]):
    ptbl.rows[0].cells[i].text = h
    for par in ptbl.rows[0].cells[i].paragraphs:
        for run in par.runs:
            run.bold = True
for c in cs:
    row = ptbl.add_row().cells
    row[0].text = c["country"]
    row[1].text = c["currency"]
    row[2].text = f"{sym[c['currency']]}{c['highest_payable_cents']/100:.2f}/mes"
    row[3].text = f"{sym[c['currency']]}{c['optimal_cents']/100:.2f}/mes"

direct = [d for d in os.listdir(GABOR)
          if os.path.isdir(os.path.join(GABOR, d)) and d not in ("mx", "br", "us")]
direct.sort(key=lambda d: os.path.getmtime(os.path.join(GABOR, d)))
if direct:
    gr = json.load(open(os.path.join(GABOR, direct[-1], "report.json"), encoding="utf-8"))
    gw = gr["numerical"]["wtp"]
    gbest = max(gw["curve"], key=lambda p: p["gross_per_exposed_cents"])
    doc.add_paragraph(
        f"Referencia global+LATAM (piloto general): precio óptimo "
        f"${gbest['price_cents']/100:.2f}/mes.")


def set_cell(cell, text, size=8, bold=False):
    cell.text = ""
    par = cell.paragraphs[0]
    run = par.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold


def set_issues(cell, issues_str, size=7.5):
    items = [i.strip() for i in (issues_str or "").split("|") if i.strip()]
    cell.text = ""
    if not items:
        return
    par = cell.paragraphs[0]
    run = par.add_run("• " + items[0])
    run.font.size = Pt(size)
    for it in items[1:]:
        par = cell.add_paragraph()
        run = par.add_run("• " + it)
        run.font.size = Pt(size)


doc.add_heading(f"Detalle por pantalla ({len(screens)})", level=1)
tbl = doc.add_table(rows=1, cols=6)
tbl.style = "Table Grid"
headers = ["Claridad", "Pantalla", "Tipo", "CTAs", "Precio", "Issues"]
for i, h in enumerate(headers):
    set_cell(tbl.rows[0].cells[i], h, size=9, bold=True)

for s in screens:
    row = tbl.add_row().cells
    clarity = s.get("clarity_0_10")
    clarity = str(clarity) if isinstance(clarity, (int, float)) else "—"
    set_cell(row[0], clarity)
    set_cell(row[1], s.get("name", ""))
    set_cell(row[2], (s.get("screenshot_kind") or "")[:40])
    set_cell(row[3], (s.get("ctas") or "")[:140])
    set_cell(row[4], (s.get("price_mentions") or "")[:70])
    set_issues(row[5], s.get("issues", ""))

out = os.path.join(AUDIT, "Mantralia-Auditoria-UX.docx")
doc.save(out)
print("DOCX:", out)
print("rows:", len(screens), "| swarm summary chars:", len(syn_summary))
